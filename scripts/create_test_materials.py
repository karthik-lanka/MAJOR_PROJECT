import os
from pathlib import Path
import docx
import fitz  # PyMuPDF
from PIL import Image, ImageDraw, ImageFont

output_dir = Path(r"c:\Users\manid\Downloads\MAJOR_PROJECT\sample_study_materials")
output_dir.mkdir(parents=True, exist_ok=True)

# 1. Plain Text: Operating Systems Fundamentals
os_txt_path = output_dir / "operating_systems_fundamentals.txt"
os_content = """# Operating Systems & Concurrency Architecture

## 1. Process Scheduling Algorithms
CPU scheduling is the basis of multi-programmed operating systems. By switching the CPU among processes, the operating system can make the computer more productive.
- First-Come, First-Served (FCFS): Simplest scheduling algorithm. Executes queued tasks sequentially. Suffers from the Convoy Effect.
- Shortest Job First (SJF): Selects the waiting process with the smallest execution time. Optimal for minimizing average turnaround time, but prone to starvation for long processes.
- Round Robin (RR): Allocates a fixed time slice or quantum (typically 10-100ms) to each process. Preemptive and fair for interactive systems.

## 2. Deadlock & Banker's Algorithm
A deadlock occurs when a set of processes are blocked because each process is holding a resource and waiting for another resource acquired by some other process.
The four Coffman conditions necessary for deadlock are:
1. Mutual Exclusion: At least one resource must be held in a non-shareable mode.
2. Hold and Wait: A process must hold at least one resource and wait for additional resources.
3. No Preemption: Resources cannot be forcibly seized from a process holding them.
4. Circular Wait: A closed loop of processes exists where each process waits for a resource held by the next.

Banker's Algorithm: Developed by Edsger Dijkstra, it tests for safety by simulating the allocation for predetermined maximum possible amounts of all resources, and decides whether allocation should be allowed to avoid unsafe states.

## 3. Virtual Memory & Paging
Virtual memory separates user logical memory from physical memory. Paging maps virtual addresses to physical frames via the Page Table.
- Translation Lookaside Buffer (TLB): A hardware associative cache that speeds up virtual-to-physical address translation.
- Thrashing: When a system spends more time paging than executing instructions, caused by insufficient allocated physical frames.
"""
os_txt_path.write_text(os_content, encoding="utf-8")
print(f"Created: {os_txt_path}")

# 2. DOCX: Software Engineering Lifecycle
docx_path = output_dir / "software_engineering_lifecycle.docx"
doc = docx.Document()
doc.add_heading("Software Engineering & Distributed Systems", level=0)

doc.add_heading("1. Agile vs Waterfall Methodology", level=1)
doc.add_paragraph(
    "Waterfall is a linear sequential development model where phases (Requirements, Design, Implementation, "
    "Verification, Maintenance) do not overlap. It is best suited for projects with fixed, unchanging specifications."
)
doc.add_paragraph(
    "Agile Methodology prioritizes adaptive planning, iterative delivery, and rapid response to change. "
    "Scrum breaks work into fixed iterations known as Sprints (typically 2 to 4 weeks), monitored via daily stand-ups, "
    "Sprint Planning, and Sprint Retrospectives."
)

doc.add_heading("2. Test-Driven Development (TDD) & CI/CD", level=1)
doc.add_paragraph(
    "Test-Driven Development follows the Red-Green-Refactor cycle:\n"
    "1. Red: Write a failing automated unit test for the desired behavior.\n"
    "2. Green: Write minimal code required to make the test pass.\n"
    "3. Refactor: Clean up and optimize the implementation while keeping tests green."
)
doc.add_paragraph(
    "Continuous Integration and Continuous Deployment (CI/CD) automates build, test, and release pipelines. "
    "Every pull request triggers automated linting, unit tests, and integration test suites before automated deployment."
)

doc.add_heading("3. Microservices vs Monolithic Architecture", level=1)
p = doc.add_paragraph()
p.add_run("Monolithic Architectures ").bold = True
p.add_run("combine database access, business logic, and UI in a single codebase. Easy to develop early, but harder to scale independently.\n")
p.add_run("Microservices Architectures ").bold = True
p.add_run("decompose systems into autonomous, loosely coupled services communicating via REST APIs or gRPC, enabling independent scaling and language-agnostic service design.")

doc.save(docx_path)
print(f"Created: {docx_path}")

# 3. PDF: Database Management Systems
pdf_path = output_dir / "database_systems_and_indexing.pdf"
doc_pdf = fitz.open()

# Page 1
page1 = doc_pdf.new_page(width=595, height=842) # A4
p1_text = """DATABASE MANAGEMENT SYSTEMS: NORMALIZATION & ACID

Chapter 1: Relational Database Normalization
Database normalization is the process of organizing relational tables to reduce data redundancy and improve data integrity.

1. First Normal Form (1NF):
- Each column must contain atomic (indivisible) values.
- No repeating groups or arrays within a record.
- Every table must possess a unique primary key.

2. Second Normal Form (2NF):
- Must satisfy all 1NF rules.
- Eliminates Partial Dependency: All non-key attributes must depend fully on the primary key, not merely a subset of a composite primary key.

3. Third Normal Form (3NF):
- Must satisfy all 2NF rules.
- Eliminates Transitive Dependency: Non-key attributes must depend directly on the primary key, and not on other non-key attributes (X -> Y and Y -> Z where Z is non-key).

4. Boyce-Codd Normal Form (BCNF):
- A stricter 3.5NF variation. For every functional dependency X -> Y, X must be a super key.
"""
page1.insert_textbox(fitz.Rect(50, 50, 545, 790), p1_text, fontsize=11, fontname="helv")

# Page 2
page2 = doc_pdf.new_page(width=595, height=842)
p2_text = """Chapter 2: Transaction Concurrency & ACID Properties

ACID guarantees safe execution of database transactions:
- Atomicity: 'All or nothing'. Either all database operations in a transaction succeed and commit, or the entire transaction is rolled back to its initial state upon failure.
- Consistency: Ensures transactions transition the database from one valid state to another, respecting all schema constraints, cascades, and foreign keys.
- Isolation: Ensures concurrent execution of transactions leaves the database in the same state as if they were executed sequentially. Standard ANSI SQL isolation levels include:
  1. Read Uncommitted
  2. Read Committed
  3. Repeatable Read
  4. Serializable
- Durability: Once a transaction commits, its effects persist permanently in non-volatile storage, even across sudden power failures or system crashes (guaranteed by Write-Ahead Logging / WAL).

Chapter 3: B+ Tree Indexing
B+ Trees are self-balancing search trees that store all actual data records in leaf nodes connected by a bidirectional linked list, while internal nodes store index navigation keys.
They maintain O(log N) lookup, insertion, and deletion complexity and provide exceptional range query performance.
"""
page2.insert_textbox(fitz.Rect(50, 50, 545, 790), p2_text, fontsize=11, fontname="helv")

doc_pdf.save(str(pdf_path))
doc_pdf.close()
print(f"Created: {pdf_path}")

# 4. PNG: Machine Learning Architecture Cheatsheet
png_path = output_dir / "machine_learning_cheatsheet.png"
img = Image.new("RGB", (900, 600), color=(15, 23, 42)) # Deep slate background
draw = ImageDraw.Draw(img)

# Draw Title
draw.text((40, 30), "MACHINE LEARNING: PARADIGMS & EVALUATION", fill=(248, 250, 252))

# Block 1: Paradigms
draw.rectangle([(40, 70), (860, 220)], outline=(59, 130, 246), width=2)
draw.text((60, 85), "1. MACHINE LEARNING PARADIGMS", fill=(96, 165, 250))
draw.text((60, 115), "- Supervised Learning: Trained on labeled dataset (X -> Y). Examples: Linear Regression, SVM, Random Forest.", fill=(226, 232, 240))
draw.text((60, 145), "- Unsupervised Learning: Discovers hidden patterns in unlabeled data. Examples: K-Means, PCA, Autoencoders.", fill=(226, 232, 240))
draw.text((60, 175), "- Reinforcement Learning: Agent learns via trial and error maximizing cumulative reward in an environment.", fill=(226, 232, 240))

# Block 2: Metrics
draw.rectangle([(40, 240), (860, 410)], outline=(16, 185, 129), width=2)
draw.text((60, 255), "2. CLASSIFICATION METRICS", fill=(52, 211, 153))
draw.text((60, 285), "- Accuracy = (TP + TN) / (TP + TN + FP + FN)", fill=(226, 232, 240))
draw.text((60, 315), "- Precision = TP / (TP + FP)  -> Crucial when False Positives are expensive (e.g., spam detection)", fill=(226, 232, 240))
draw.text((60, 345), "- Recall = TP / (TP + FN)     -> Crucial when False Negatives are dangerous (e.g., disease diagnosis)", fill=(226, 232, 240))
draw.text((60, 375), "- F1-Score = 2 * (Precision * Recall) / (Precision + Recall) -> Harmonic mean balancing precision and recall", fill=(226, 232, 240))

# Block 3: Optimization
draw.rectangle([(40, 430), (860, 560)], outline=(245, 158, 11), width=2)
draw.text((60, 445), "3. GRADIENT DESCENT OPTIMIZATION", fill=(251, 191, 36))
draw.text((60, 475), "- Objective: Minimize loss function J(theta) by iteratively updating parameters opposite to the gradient.", fill=(226, 232, 240))
draw.text((60, 505), "- Update Rule: theta_new = theta_old - (learning_rate * grad(J))", fill=(226, 232, 240))
draw.text((60, 535), "- Learning rate too high causes divergence; learning rate too low causes slow convergence.", fill=(226, 232, 240))

img.save(png_path)
print(f"Created: {png_path}")
print("All 4 multi-format test materials created successfully!")
